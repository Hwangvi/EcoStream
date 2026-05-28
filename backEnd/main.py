import os
import io
import json
from app.utils.kafka_helper import crear_productor
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import VerticalBarChart
from datetime import datetime, timedelta
from typing import List, Optional
from pydantic import BaseModel
from fastapi import FastAPI, HTTPException, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openai import OpenAI
from dotenv import load_dotenv
from app.database.database import EcoStreamRepository, get_db_connection
from contextlib import asynccontextmanager
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import configurar_cors

load_dotenv()

@asynccontextmanager
async def lifespan(app: FastAPI):
    print("🧹 Ejecutando limpieza inicial de la base de datos...")
    try:
        EcoStreamRepository.purgar_datos_antiguos()
        print("✅ Sistema listo y saneado.")
    except Exception as e:
        print(f"⚠️ Error al limpiar en el inicio: {e}")
    yield
    print("🔌 Apagando EcoStream Engine...")

app = FastAPI(
    title="EcoStream Analytics Engine",
    description="Core analítico de telemetría urbana y auditoría con IA",
    lifespan=lifespan
)
configurar_cors(app)

origins = os.getenv("CORS_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173").split(",")
app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class PeticionAnalisis(BaseModel):
    zona: str

@app.get("/api/alertas/exportar")
async def exportar_registro_pdf():
    try:
        filas = EcoStreamRepository.obtener_ultimas_100_filas()

        if not filas:
            return {"status": "error", "message": "No se encontraron registros en este rango de tiempo."}
        
        if len(filas) > 500:
            filas = filas[:500]
            
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(letter),title="Reporte de Auditoría - EcoStream",
    author="HwangVi",
    subject="Reporte técnico de alertas y mediciones")
        elements = []
        styles = getSampleStyleSheet()

        elements.append(Paragraph("Reporte de Auditoría: Últimos 100 registros", styles['Title']))
        elements.append(Spacer(1, 12))

        data = [["Tipo", "Sensor", "Valor", "Descripción", "Fecha/Hora"]]
        
        for f in filas:
            tipo = f.get('tipo') or 'Desconocido'
            codigo = f.get('codigo_sensor') or 'S/C'
            valor = f.get('valor') or 0.0
            desc = f.get('descripcion') or 'Sin descripción'
            fecha = str(f.get('fecha_evento'))[:16] if f.get('fecha_evento') else "N/A"
            
            es_alerta = tipo == 'ALERTA'
            color_texto = colors.red if es_alerta else colors.black

            fila = [
                tipo,
                Paragraph(f"<font color='{color_texto}'>{codigo}</font>", styles['Normal']),
                Paragraph(f"<font color='{color_texto}'>{str(round(float(valor), 2))}</font>", styles['Normal']),
                Paragraph(f"<font color='{color_texto}'>{desc}</font>", styles['Normal']),
                fecha
            ]
            data.append(fila)

        table = Table(data, colWidths=[60, 100, 60, 250, 100])
        table.setStyle(TableStyle([
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ]))
        
        elements.append(table)
        doc.build(elements)
        
        buffer.seek(0)
        return StreamingResponse(
            buffer, 
            media_type="application/pdf", 
            headers={"Content-Disposition": "attachment; filename=Reporte_EcoStream.pdf"}
        )
    except Exception as e:
        print(f"ERROR CRÍTICO: {e}")
        raise HTTPException(status_code=500, detail="Error en servidor")

@app.get("/api/alertas/conteo")
async def contar_alertas(rango: int = 1):
    count = EcoStreamRepository.contar_alertas(rango)
    return {"total": count}

    
@app.post("/api/alertas/inyectar-anomalia")
async def inyectar_anomalia(zona: str):
    payload = {
        "codigo_sensor": f"SEN-{zona.upper()}-CO2", 
        "valor": 999.0
    }
    
    producer = crear_productor()
    producer.send('telemetria-sensores', value=payload)
    producer.flush()
    producer.close()
    
    return {"status": "Anomalía inyectada correctamente"}

@app.post("/api/alertas/reset")
async def reset_alertas():
    try:
        EcoStreamRepository.limpiar_alertas()
        return {"status": "success", "message": "Historial de alertas limpiado"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al limpiar: {str(e)}")
    
from app.utils.kafka_helper import crear_productor

@app.post("/api/ia/analizar")
async def analizar_zona_con_ia(data: PeticionAnalisis):
    try:
        filas = EcoStreamRepository.obtener_alertas_recientes()
        
        alertas_zona = []
        for f in filas:
            alerta_dict = dict(f)
            codigo = str(alerta_dict.get("codigo_sensor", ""))
            
            if data.zona == "TODAS" or f"-{data.zona}-" in codigo:
                try: valor_num = float(alerta_dict.get("valor_disparado", 0.0))
                except: valor_num = 0.0

                alertas_zona.append({
                    "codigo": codigo,
                    "valor": valor_num,
                    "descripcion": str(alerta_dict.get("descripcion", "Anomalía territorial."))
                })

        if not alertas_zona:
            reporte_texto = f"Métricas estables en el cuadrante {data.zona}. No se registran picos críticos de contaminación."
        else:
            reporte_texto = f"--- INICIO REPORTE ANOMALÍAS - ÁREA: {data.zona} ---\n"
            for idx, a in enumerate(alertas_zona):
                reporte_texto += f"Incidente #{idx+1}: Dispositivo {a['codigo']} | Valor: {a['valor']} | Diagnóstico: {a['descripcion']}\n"
            reporte_texto += "--- FIN REPORTE ---"

        client_groq = OpenAI(
            base_url="https://api.groq.com/openai/v1",
            api_key=os.getenv("GROQ_API_KEY")
        )

        system_prompt = (
            "Eres un experto senior en Smart Cities y resiliencia medioambiental urbana. Tu tarea es auditar un log "
            "de métricas críticas y dictaminar un plan de mitigación técnica inmediato.\n\n"
            "Debes responder OBLIGATORIAMENTE con un objeto JSON estructurado que contenga exactamente estas dos llaves:\n"
            "1. 'diagnostico': Evaluación profesional del estado del cuadrante urbano (máximo 3 líneas).\n"
            "2. 'recomendaciones': Un array de strings (máximo 3) con soluciones de infraestructura concretas.\n\n"
            "Ejemplo estricto de salida:\n"
            '{"diagnostico": "Análisis...", "recomendaciones": ["Medida 1", "Medida 2"]}'
        )

        response = client_groq.chat.completions.create(
            model="llama-3.1-8b-instant",
            response_format={"type": "json_object"},  
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": f"Audita el siguiente reporte de telemetría urbana:\n\n{reporte_texto}"}
            ],
            temperature=0.2
        )

        resultado_ia = json.loads(response.choices[0].message.content)

        return {
            "status": "success",
            "zona_auditada": data.zona,
            "alertas_procesadas": len(alertas_zona),
            "diagnostico": resultado_ia.get("diagnostico", "Auditoría no concluyente."),
            "recomendaciones": resultado_ia.get("recomendaciones", [])
        }
        
    except Exception as e:
        print(f"⚠️ Fallo en Agente Cognitivo Groq: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Error analítico del LLM: {str(e)}.")


@app.get("/api/dashboard")
async def obtener_dashboard():
    try:
        return EcoStreamRepository.obtener_ultimas_mediciones()
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/alertas")
async def obtener_alertas():
    try:
        return EcoStreamRepository.obtener_alertas_recientes()
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))